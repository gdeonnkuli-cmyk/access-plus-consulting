"""Extrait les cantiques du recueil Word (.docx) vers data/cantiques.json.

Usage : python3 tools/extract_cantiques.py chemin/recueil.docx data/cantiques.json

Repères de mise en forme utilisés dans le recueil :
  - « CANTIQUE NN »                     -> début d'un cantique
  - « NN<tab>Mélodie. Réf. »            -> mélodie / référence
  - paragraphe en retrait commençant par un chiffre -> strophe
  - paragraphe en retrait sans numéro (centré)      -> refrain
  - paragraphe aligné à droite          -> auteur / traducteur
  - paragraphe centré sans retrait      -> titre de section (thème)
"""
import json
import re
import sys

import docx
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT as AL

HEAD = re.compile(r"^CANTIQUE\s+(\d+)\s*$")
TUNE = re.compile(r"^(\d+)(?:\s*\t\s*(.*))?$")
STANZA = re.compile(r"^(\d+)\s{2,}(.*)$", re.S)


def clean(s):
    s = s.replace("’", "'").replace("‘", "'")
    return "\n".join(l.strip() for l in s.strip().split("\n") if l.strip())


def main(src, dst):
    d = docx.Document(src)
    hymns, cur, section, pending_section = [], None, None, []
    seen = set()
    for t in d.tables[3:]:  # les 3 premières tables = index alphabétique
        for r in t.rows:
            for c in r.cells:
                if c._tc in seen:
                    continue
                seen.add(c._tc)
                for p in c.paragraphs:
                    txt = p.text.strip()
                    if not txt:
                        continue
                    indent = p.paragraph_format.left_indent
                    m = HEAD.match(txt)
                    if m:
                        if pending_section:
                            section = " ".join(pending_section).strip().rstrip(".").strip()
                            pending_section = []
                        cur = {"n": int(m.group(1)), "melodie": "", "section": section,
                               "auteur": "", "parties": []}
                        hymns.append(cur)
                        continue
                    if cur is None:
                        continue
                    m = TUNE.match(txt)
                    if m and int(m.group(1)) == cur["n"] and not cur["parties"]:
                        cur["melodie"] = (m.group(2) or "").strip()
                        continue
                    if indent:
                        m = STANZA.match(p.text.strip())
                        if m:
                            cur["parties"].append({"type": "strophe", "n": int(m.group(1)),
                                                   "texte": clean(m.group(2))})
                        else:
                            cur["parties"].append({"type": "refrain", "texte": clean(txt)})
                        continue
                    if p.alignment == AL.RIGHT:
                        cur["auteur"] = txt.rstrip(".").strip()
                    elif p.alignment == AL.CENTER:
                        pending_section.append(clean(txt).replace("\n", " "))
                    else:
                        # texte libre non indenté : on le garde comme refrain
                        cur["parties"].append({"type": "refrain", "texte": clean(txt)})

    hymns.sort(key=lambda h: h["n"])
    for h in hymns:
        first = next((p for p in h["parties"] if p["type"] == "strophe"), None) or \
                (h["parties"][0] if h["parties"] else {"texte": ""})
        h["titre"] = first["texte"].split("\n")[0].rstrip(" ,;:.!").strip()
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(hymns, f, ensure_ascii=False, separators=(",", ":"))

    nums = [h["n"] for h in hymns]
    dup = sorted({n for n in nums if nums.count(n) > 1})
    manquants = [n for n in range(1, max(nums) + 1) if n not in nums]
    vides = [h["n"] for h in hymns if not h["parties"]]
    print(f"{len(hymns)} cantiques | doublons {dup} | manquants {manquants} | vides {vides}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
